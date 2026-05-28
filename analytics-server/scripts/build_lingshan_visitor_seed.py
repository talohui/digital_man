#!/usr/bin/env python3
"""Build de-identified row-level visitor seed data for Lingshan dashboard.

The source workbook is already filtered to Lingshan-related attractions. This
script intentionally excludes tourist_id, user_nickname and attraction_content.
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_SOURCE = Path("/Users/MR/Desktop/软件杯/地图/outputs/lingshan_filter/灵山胜境相关景点筛选结果.xlsx")
DEFAULT_OUTPUT = Path("src/main/resources/visitor-behavior/lingshan-visitor-behavior-seed-v1.json")
SOURCE = "history_lingshan_sample"


def number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str) and value.strip():
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


def text(value: Any) -> str:
    return str(value or "").strip()


def to_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.strip():
        try:
            return datetime.fromisoformat(value.strip()).date().isoformat()
        except ValueError:
            return None
    return None


def age_band(age: float | None) -> str:
    if age is None:
        return "未知"
    if age <= 24:
        return "18-24"
    if age <= 34:
        return "25-34"
    if age <= 44:
        return "35-44"
    if age <= 59:
        return "45-59"
    return "60+"


def normalize_gender(value: Any) -> str:
    gender = text(value)
    return gender if gender in {"男", "女"} else "未知"


def round2(value: float | None) -> float:
    return round(value or 0.0, 2)


def build(source: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: pos for pos, name in enumerate(headers)}
    required = {
        "age", "gender", "attraction_name", "attraction_type", "visit_date",
        "stay_duration", "ticket_cost", "food_cost", "shopping_cost",
        "transport_cost", "entertainment_cost", "total_cost", "group_size",
        "satisfaction",
    }
    missing = sorted(required - set(index))
    if missing:
        raise SystemExit(f"Missing required columns: {missing}")

    records: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        age = number(row[index["age"]])
        satisfaction_raw = number(row[index["satisfaction"]])
        record = {
            "source": SOURCE,
            "visitorId": f"history-{row_number:04d}",
            "ticketId": f"history-ticket-{row_number:04d}",
            "attractionName": text(row[index["attraction_name"]]),
            "attractionType": text(row[index["attraction_type"]]),
            "ageBand": age_band(age),
            "gender": normalize_gender(row[index["gender"]]),
            "visitDate": to_date(row[index["visit_date"]]),
            "stayHours": round2(number(row[index["stay_duration"]])),
            "ticketCost": round2(number(row[index["ticket_cost"]])),
            "foodCost": round2(number(row[index["food_cost"]])),
            "shoppingCost": round2(number(row[index["shopping_cost"]])),
            "transportCost": round2(number(row[index["transport_cost"]])),
            "entertainmentCost": round2(number(row[index["entertainment_cost"]])),
            "totalCost": round2(number(row[index["total_cost"]])),
            "groupSize": int(number(row[index["group_size"]]) or 1),
            "satisfaction": round2(satisfaction_raw * 20.0 if satisfaction_raw is not None else None),
        }
        records.append(record)
    return records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    records = build(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(records)} Lingshan visitor records to {args.output}")


if __name__ == "__main__":
    main()
